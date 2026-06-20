"""Document upload flow helpers — aligned with legacy PromptEnhancement behavior."""

from __future__ import annotations

import asyncio
import io
import re
from pathlib import Path
from typing import Any, Callable

DOCUMENT_FLOW_ALLOWED_EXTENSIONS = [
    ".pdf",
    ".pptx",
    ".docx",
    ".txt",
    ".csv",
    ".xlsx",
    ".xls",
    ".jpg",
    ".jpeg",
    ".png",
    ".webp",
    ".bmp",
    ".tiff",
    ".gif",
    ".md",
]

MAX_DOCUMENT_SIZE_MB = 50


def validate_file_upload(
    file_content: bytes,
    filename: str,
    *,
    allowed_extensions: list[str] | None = None,
    max_size_mb: int = MAX_DOCUMENT_SIZE_MB,
) -> None:
    """Validate upload size and extension (legacy-compatible)."""
    if not filename:
        raise ValueError("No file provided")
    ext = Path(filename).suffix.lower()
    allowed = allowed_extensions or DOCUMENT_FLOW_ALLOWED_EXTENSIONS
    if ext not in allowed:
        raise ValueError(f"Unsupported file type: {ext or '(none)'}")
    max_bytes = max_size_mb * 1024 * 1024
    if len(file_content) > max_bytes:
        raise ValueError(f"File exceeds maximum size of {max_size_mb} MB")


def chunk_text(text: str, chunk_size: int = 350, overlap: int = 100) -> list[str]:
    """Split text into overlapping word-based chunks (legacy-compatible)."""
    words = text.split()
    if not words:
        return []
    step = max(1, chunk_size - overlap)
    chunks: list[str] = []
    for i in range(0, len(words), step):
        chunk = " ".join(words[i : i + chunk_size]).strip()
        if chunk:
            chunks.append(chunk)
    return chunks


def generate_document_summary(text: str, max_sentences: int = 8) -> str:
    """Build a whole-document extractive summary (ported from PromptEnhancement)."""
    raw = (text or "").strip()
    if not raw:
        return ""

    normalized = re.sub(r"[ \t]+", " ", raw)
    normalized = re.sub(r"\n{2,}", "\n", normalized).strip()

    sentence_candidates = [
        s.strip(" -•\t\n")
        for s in re.split(r"(?<=[.!?])\s+|\n", normalized)
        if s and s.strip()
    ]
    if not sentence_candidates:
        return re.sub(r"\s+", " ", raw).strip()

    filtered: list[str] = []
    for sentence in sentence_candidates:
        low = sentence.lower().strip()
        if len(sentence) < 35:
            continue
        if re.match(r"^(for:|version\b|classification\b|page\s+\d+|part\s+[ivx]+)", low):
            continue
        if re.match(r"^\d+(\.\d+)?\s+[a-z]", low):
            continue
        filtered.append(sentence)

    sentences = filtered or sentence_candidates
    total_sentences = len(sentences)

    stop_words = {
        "the", "a", "an", "and", "or", "but", "if", "then", "this", "that", "these", "those",
        "is", "are", "was", "were", "be", "been", "being", "to", "of", "in", "on", "for", "with",
        "as", "at", "by", "from", "it", "its", "your", "you", "we", "our", "they", "their", "them",
    }

    token_re = re.compile(r"[a-zA-Z][a-zA-Z0-9_-]{2,}")
    freq: dict[str, int] = {}
    tokenized_sentences: list[list[str]] = []
    for sentence in sentences:
        tokens = [t.lower() for t in token_re.findall(sentence) if t.lower() not in stop_words]
        tokenized_sentences.append(tokens)
        for token in tokens:
            freq[token] = freq.get(token, 0) + 1

    scored: list[tuple[float, int, str]] = []
    for idx, sentence in enumerate(sentences):
        tokens = tokenized_sentences[idx]
        if not tokens:
            continue
        lexical_score = sum(freq.get(token, 0) for token in tokens) / max(1, len(tokens))
        position_ratio = idx / max(1, total_sentences - 1)
        if position_ratio <= 0.2:
            position_bonus = 1.12
        elif position_ratio >= 0.8:
            position_bonus = 1.08
        else:
            position_bonus = 1.0
        structural_bonus = 1.0
        if any(
            keyword in sentence.lower()
            for keyword in ("must", "required", "non-negotiable", "process", "protocol", "blocker", "accountable")
        ):
            structural_bonus += 0.08
        scored.append((lexical_score * position_bonus * structural_bonus, idx, sentence))

    if not scored:
        return re.sub(r"\s+", " ", raw).strip()

    target = max(4, min(max_sentences, 10))
    bucket_count = min(4, max(1, total_sentences))
    bucket_size = max(1, total_sentences // bucket_count)
    selected_indices: set[int] = set()
    for bucket in range(bucket_count):
        start = bucket * bucket_size
        end = total_sentences if bucket == bucket_count - 1 else min(total_sentences, (bucket + 1) * bucket_size)
        bucket_items = [item for item in scored if start <= item[1] < end]
        if bucket_items:
            selected_indices.add(max(bucket_items, key=lambda x: x[0])[1])

    for _, idx, _ in sorted(scored, key=lambda x: x[0], reverse=True):
        if len(selected_indices) >= target:
            break
        selected_indices.add(idx)

    final_lines: list[str] = []
    seen: set[str] = set()
    for line in [sentences[i] for i in sorted(selected_indices)]:
        key = re.sub(r"[^a-z0-9 ]", "", line.lower()).strip()
        if key and key not in seen:
            seen.add(key)
            final_lines.append(line)

    summary = " ".join(final_lines).strip()
    return re.sub(r"\s+", " ", summary)


def _extract_txt(file_content: bytes) -> str:
    try:
        return file_content.decode("utf-8").strip()
    except UnicodeDecodeError:
        try:
            return file_content.decode("latin-1").strip()
        except UnicodeDecodeError:
            return file_content.decode("utf-8", errors="ignore").strip()


def _extract_pdf(file_content: bytes) -> str:
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=file_content, filetype="pdf")
        text = "".join(page.get_text() for page in doc)
        doc.close()
        return text.strip()
    except ImportError:
        try:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                parts = [page.extract_text() or "" for page in pdf.pages]
            return "\n".join(parts).strip()
        except ImportError:
            try:
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(file_content))
                parts = [page.extract_text() or "" for page in reader.pages]
                return "\n".join(parts).strip()
            except ImportError as exc:
                raise ValueError(
                    "PDF parsing library not available. Install PyMuPDF, pdfplumber, or pypdf."
                ) from exc


def _extract_docx(file_content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("python-docx is required for DOCX uploads") from exc
    doc = Document(io.BytesIO(file_content))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs).strip()


def _extract_pptx(file_content: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ValueError("python-pptx is required for PPTX uploads") from exc
    prs = Presentation(io.BytesIO(file_content))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts).strip()


def _extract_excel(file_content: bytes) -> str:
    try:
        import pandas as pd
    except ImportError as exc:
        raise ValueError("pandas/openpyxl is required for spreadsheet uploads") from exc

    sheets = pd.read_excel(io.BytesIO(file_content), sheet_name=None)
    rendered: list[str] = []
    for sheet_name, frame in sheets.items():
        rendered.append(f"Sheet: {sheet_name}")
        rendered.append(frame.to_csv(index=False))
    return "\n".join(rendered).strip()


def _extract_image_caption(file_content: bytes, filename: str) -> str:
    # Best-effort caption without vision model — filename + size metadata.
    return f"Image upload: {filename} ({len(file_content)} bytes)."


_EXTRACTORS: dict[str, Callable[[bytes], str]] = {
    ".txt": _extract_txt,
    ".md": _extract_txt,
    ".csv": _extract_txt,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_excel,
    ".xls": _extract_excel,
}

_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff", ".gif"}


async def extract_text_for_document_flow(file_bytes: bytes, filename: str) -> dict[str, Any]:
    """Extract text/caption from supported file types (legacy-compatible)."""
    ext = Path(filename).suffix.lower()
    if ext in _IMAGE_EXTENSIONS:
        text = await asyncio.to_thread(_extract_image_caption, file_bytes, filename)
        content_type = "image"
    else:
        extractor = _EXTRACTORS.get(ext)
        if extractor is None:
            raise ValueError(f"Unsupported file type for document flow: {ext}")
        text = await asyncio.to_thread(extractor, file_bytes)
        content_type = "document"

    text = (text or "").strip()
    if not text:
        raise ValueError("Could not extract content from file")

    return {
        "text": text,
        "content_type": content_type,
        "file_extension": ext,
    }
